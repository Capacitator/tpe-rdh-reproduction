#!/usr/bin/env python3
"""Copy the corrected legacy-format report and append actual-r1 NIST results.

The input report is read-only; output uses a new filename and will not overwrite.
"""
from __future__ import annotations
import csv
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
SOURCE = REPO / "authenticated_tpe_nist_v2" / "output" / "report" / "authenticated_tpe_nist_v2_fidelity_correction.docx"
TARGET = REPO / "authenticated_tpe_nist_v2" / "output" / "report" / "authenticated_tpe_final_corrected_with_actual_r1.docx"
SUMMARY = ROOT / "output" / "summary.md"
TESTS = ROOT / "output" / "test_by_test_summary.csv"


def csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def add_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(8)


def mark_header_row(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_heading(document, text: str, size: int = 12):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return paragraph


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    if TARGET.exists():
        raise FileExistsError(f"Refusing to overwrite existing report: {TARGET}")
    summary = SUMMARY.read_text(encoding="utf-8")
    rows = csv_rows(TESTS)
    if len(rows) != 15:
        raise AssertionError(f"expected 15 named NIST test rows; got {len(rows)}")

    doc = Document(str(SOURCE))
    doc.add_page_break()
    add_heading(doc, "Supplement: NIST assessment of actual block-r1 outputs", 14)
    p = doc.add_paragraph(
        "This new appendix adds a direct NIST STS test of the two-bit r1 values consumed by "
        "the encryption transform. It supplements—does not replace—the earlier `block_r1` "
        "category, which tested full HMAC-SHA256 digests. No cipher parameter or identifier "
        "was changed to target p-values, and no stream was removed."
    )
    p.paragraph_format.space_after = Pt(6)
    add_heading(doc, "Protocol and overall results", 11)
    for text in [
        "Official NIST STS 2.1.2; all default tests; 100 independent streams, 1,000,000 bits per stream; first-level α=0.01; second-level uniformity threshold 0.0001.",
        "Actual r1 definition: int.from_bytes(HMAC-SHA256(Kr1_c, b'r1' || u16be(block_id))[:4], 'big') % 4. Each 2-bit result is packed MSB-first in channel/block order. Five hundred thousand r1 symbols form each 1,000,000-bit input stream.",
        "The 100 streams use disjoint predeclared deterministic UserKey indices 10,000–75,199 (65,200 key instances total). The binary stream hashes are all unique; no identifier from the original campaign is altered or reused.",
        "First-level entries across 188 NIST components: 17,630 pass, 156 fail, and 1,014 N/A; 17,786 applicable p-values, 99.123% pass. N/A entries are NIST Random Excursions applicability outcomes (39 streams for each excursion family), not dropped failures.",
        "Second-level: 188/188 valid uniformity p-values are at least 0.0001; 0 fail the NIST uniformity threshold. Minimum second-level p = 0.000199128574335 (NonOverlappingTemplate), disclosed because it is near the threshold. Six template-component uniformity values are below 0.01 but remain above 0.0001.",
        "The approximate NIST three-sigma pass-proportion check flagged 0/188 components; exact one-sided binomial checks with Holm correction at 0.01 flagged 0/188. First-level failures are still retained and reported below.",
    ]:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.font.size = Pt(9)

    add_heading(doc, "Results by named NIST test", 11)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.45), Inches(1.00), Inches(0.95), Inches(1.75), Inches(1.20)]
    header = ["NIST test", "1st P/F/N-A", "Pass rate", "2nd-level U: n / <1e-4 / <0.01", "Min U p"]
    for i, text in enumerate(header):
        table.columns[i].width = widths[i]
        add_cell(table.rows[0].cells[i], text, True)
    mark_header_row(table.rows[0])
    for row in rows:
        p = int(row["first_level_pass"])
        f = int(row["first_level_fail"])
        na = int(row["first_level_NA_component_values"])
        display_names = {
            "BlockFrequency": "Block frequency",
            "CumulativeSums": "Cumulative sums",
            "LongestRun": "Longest run",
            "NonOverlappingTemplate": "Non-overlapping template",
            "OverlappingTemplate": "Overlapping template",
            "ApproximateEntropy": "Approximate entropy",
            "RandomExcursions": "Random excursions",
            "RandomExcursionsVariant": "Random excursions variant",
            "LinearComplexity": "Linear complexity",
        }
        second = (f"{row['second_level_values']} / "
                  f"{row['second_level_failures_at_0_0001']} / "
                  f"{row['second_level_values_below_0_01_exploratory']}")
        values = [display_names.get(row["nist_test"], row["nist_test"]), f"{p}/{f}/{na}",
                  f"{100*float(row['first_level_pass_rate']):.3f}%",
                  second, row["minimum_second_level_p"] or "N/A"]
        cells = table.add_row().cells
        for i, text in enumerate(values):
            cells[i].width = widths[i]
            add_cell(cells[i], text)

    doc.add_paragraph(
        "The 1st-level counts are per test component (for example, NonOverlappingTemplate has "
        "148 components). The complete first-level p-values, complete second-level uniformity "
        "values, N/A entries, raw reports, stream hashes and artifact hashes are retained in the "
        "supplemental experiment output. In the second-level column, n / <1e-4 / <0.01 means "
        "number of uniformity p-values / count below NIST's official threshold / count below 0.01. "
        "NIST statistical tests do not establish cryptographic security."
    )
    add_heading(doc, "Interpretation", 11)
    doc.add_paragraph(
        "This is a genuine input-selection correction: the actual reduced r1 values are now tested "
        "in addition to the raw HMAC digests. The observed Frequency (100/100) and Runs (100/100) "
        "results are good for this run, but the algorithm was not altered to improve p-values. "
        "The other first-level failures remain visible in the table, and no claim is made that every "
        "individual test passes. The prior corrected image-fidelity results and all earlier Word "
        "reports are preserved unchanged."
    )
    doc.save(TARGET)
    print(f"Created new report: {TARGET}")
    print(f"Source report preserved: {SOURCE}")

if __name__ == "__main__":
    main()
