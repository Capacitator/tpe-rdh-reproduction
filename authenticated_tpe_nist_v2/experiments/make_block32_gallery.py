#!/usr/bin/env python3
"""Create a professor-facing b=32 gallery matching the earlier DOCX layout."""
from __future__ import annotations
import csv
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
IMG=ROOT/"output"/"metrics"/"block32_corrected_images"
MANIFEST=IMG/"manifest.csv"
TARGET=IMG/"professor_32x32_block_effect_gallery.docx"


def main():
    with MANIFEST.open(newline="",encoding="utf-8") as f: rows=list(csv.DictReader(f))
    if len(rows)!=6: raise RuntimeError(f"Expected six validated image pairs; got {len(rows)}")
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.45);sec.bottom_margin=Inches(.45)
    sec.left_margin=Inches(.55);sec.right_margin=Inches(.55)
    normal=doc.styles["Normal"];normal.font.name="Arial";normal.font.size=Pt(9)
    title=doc.add_heading("32×32 Block-Encrypted Image Examples",0)
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    intro=doc.add_paragraph(
        "Each right-hand panel is a newly generated encrypted output using the previous TPE/RDH "
        "pipeline with block size b=32. The images are shown at native input dimensions, scaled only "
        "to fit this page. All six runs verified exact recovery, payload recovery, and preservation "
        "of each 32×32 per-channel block sum."
    )
    intro.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        heading=doc.add_heading(f"{row['image'].capitalize()} — original vs. b=32 encrypted image",level=2)
        heading.paragraph_format.keep_with_next=True
        detail=doc.add_paragraph(
            f"Block size: 32×32 | Exact recovery: {row['exact_recovery']} | "
            f"Payload recovered: {row['payload_recovered']} | "
            f"Block sums preserved: {row['thumbnail_block_sums_preserved']}"
        )
        detail.paragraph_format.keep_with_next=True
        table=doc.add_table(rows=2,cols=2)
        table.autofit=False
        table.columns[0].width=Inches(3.55);table.columns[1].width=Inches(3.55)
        for col,label in enumerate(("Original input","32×32 block-encrypted output")):
            cell=table.cell(0,col);cell.text=label
            cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
            pic=table.cell(1,col).paragraphs[0];pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
            filename=row["original_png"] if col==0 else row["encrypted_png"]
            pic.add_run().add_picture(str(IMG/filename),width=Inches(3.20))
    note=doc.add_paragraph(
        "Method note: These images are from the separate chaotic TPE/RDH implementation, not the "
        "authenticated block-group prototype. The authenticated method processes 32×32 blocks but "
        "does not spatially permute block positions; it is therefore not expected to show this same "
        "block-wise appearance. No visual overlay or post-processing has been applied."
    )
    note.paragraph_format.space_before=Pt(5)
    doc.save(TARGET)
    print(f"Created {TARGET}; {len(rows)} original/encrypted pairs included.")

if __name__=="__main__":main()
