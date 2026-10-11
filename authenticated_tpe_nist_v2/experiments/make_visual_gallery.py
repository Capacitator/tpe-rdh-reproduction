#!/usr/bin/env python3
"""Export exact original/final-marked image pairs and a professor-ready DOCX gallery.

Uses the same image pool, UserKey and ImageID as run_metrics. Each saved pixel
array is verified against the SHA-256 values in image_quality.csv.
"""
from __future__ import annotations
import csv
import hashlib
import sys
from pathlib import Path

import numpy as np
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import atpe_v2 as v2  # noqa: E402
import streams as S  # noqa: E402

OUT = ROOT / "output" / "metrics" / "visual_gallery"
METRICS = ROOT / "output" / "metrics" / "image_quality.csv"


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    with METRICS.open(newline="", encoding="utf-8") as f:
        metric_rows = list(csv.DictReader(f))
    if len(metric_rows) != len(S.UCT_NAMES):
        raise RuntimeError(f"Expected six UCT rows; found {len(metric_rows)}")

    image_records = []
    for index, row in enumerate(metric_rows):
        label = row["image"]
        original = np.ascontiguousarray(S.load_image(index), dtype=np.uint8)
        key = S.pool_key(5000 + index)
        image_id = S.pool_image_id(5000 + index)
        protected = v2.protect_image(original, key, image_id)
        marked = np.ascontiguousarray(protected.marked_image, dtype=np.uint8)
        verified = v2.verify_and_decrypt(marked, key, image_id)
        original_bytes = original.tobytes(order="C")
        marked_bytes = marked.tobytes(order="C")
        original_hash = sha256(original_bytes)
        marked_hash = sha256(marked_bytes)
        if original_hash != row["image_sha256"]:
            raise AssertionError(f"input image does not match metrics row: {label}")
        if marked_hash != row["ciphertext_sha256"]:
            raise AssertionError(f"marked ciphertext does not match metrics row: {label}")
        if not verified.accepted or not np.array_equal(verified.recovered_image, original):
            raise AssertionError(f"marked ciphertext did not authenticate/recover exactly: {label}")

        original_path = OUT / f"{label}_original.png"
        marked_path = OUT / f"{label}_encrypted_marked.png"
        Image.fromarray(original, mode="RGB").save(original_path, format="PNG", optimize=True)
        Image.fromarray(marked, mode="RGB").save(marked_path, format="PNG", optimize=True)
        image_records.append({
            "image": label,
            "mode": protected.mode,
            "authentication_accepted": verified.accepted,
            "exact_recovery": bool(np.array_equal(verified.recovered_image, original)),
            "input_pixel_sha256": original_hash,
            "marked_pixel_sha256": marked_hash,
            "input_png_sha256": sha256(original_path.read_bytes()),
            "marked_png_sha256": sha256(marked_path.read_bytes()),
            "input_vs_final_marked_psnr_db": row["input_vs_encrypted_psnr_db"],
            "input_vs_final_marked_ssim": row["input_vs_encrypted_ssim"],
            "original_png": original_path.name,
            "final_authentication_marked_png": marked_path.name,
        })

    with (OUT / "image_manifest.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(image_records[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(image_records)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    title = doc.add_heading("Visual encrypted-image examples", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro = doc.add_paragraph(
        "Each right-hand image is the final encrypted image after authentication information "
        "has been embedded. It is generated using the same image, key, ImageID and final "
        "ciphertext verified in the corrected PSNR evaluation. Images are shown at native "
        "512 × 512 RGB pixels (scaled only to fit this page)."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, record in enumerate(image_records):
        heading = doc.add_heading(
            f"{record['image'].capitalize()} — input vs. final marked ciphertext", level=2)
        heading.paragraph_format.keep_with_next = True
        psnr = float(record["input_vs_final_marked_psnr_db"])
        ssim = float(record["input_vs_final_marked_ssim"])
        metric = doc.add_paragraph(
            f"Input-to-final PSNR: {psnr:.4f} dB | SSIM: {ssim:.4f} | "
            f"Mode: {record['mode']} | Authentication accepted; exact recovery verified."
        )
        metric.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=2, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(3.55)
        table.columns[1].width = Inches(3.55)
        for col, label in enumerate(("Original input", "Final encrypted + authentication data")):
            cell = table.cell(0, col)
            cell.text = label
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph = table.cell(1, col).paragraphs[0]
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            filename = record["original_png"] if col == 0 else record["final_authentication_marked_png"]
            picture_paragraph.add_run().add_picture(str(OUT / filename), width=Inches(3.25))

    doc.add_paragraph(
        "Measurement note: PSNR/SSIM compare the original input to the final authentication-marked "
        "ciphertext. The previously reported ~27.77 dB Baboon quantity compared the marked output "
        "to the Step-2 image before tag embedding and therefore measured marking-only distortion, "
        "not original-to-encrypted visual fidelity."
    )
    docx_path = OUT / "professor_encrypted_image_gallery.docx"
    doc.save(docx_path)
    print(f"Saved six verified original/final pairs and gallery: {docx_path}")
    print(f"Pixel hashes match corrected metrics: {len(image_records)}/6; all authenticate and recover exactly.")


if __name__ == "__main__":
    main()
