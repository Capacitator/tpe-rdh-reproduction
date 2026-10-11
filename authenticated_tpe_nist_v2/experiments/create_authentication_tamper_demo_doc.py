#!/usr/bin/env python3
"""Build a concise two-page Word demo with full images and exact tamper results."""
from __future__ import annotations
import csv, hashlib, sys
from pathlib import Path
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]; WORK=ROOT.parent
sys.path.insert(0,str(ROOT/'src'));sys.path.insert(0,str(WORK/'tpe_rdh_reproduction'/'src'))
import block32_exact_order_authenticated_tpe as exact
import block32_authenticated_tpe as common
BASE=ROOT/'output'/'metrics'/'flowchart_exact_order_compensated'
DEMO=BASE/'cross_image_splice_test'; BLOCK=32
KEY=hashlib.sha256(b'public reproducible research authentication key fixture').digest()
BLUE='1F4E79'

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),fill);tcPr.append(s)
def margins(cell):
    tcPr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for side,val in [('top',50),('start',50),('bottom',50),('end',50)]:
        e=OxmlElement(f'w:{side}');e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');m.append(e)
    tcPr.append(m)
def style_run(r,size=9,bold=False,color=None,italic=False):
    r.font.name='Aptos';r.font.size=Pt(size);r.bold=bold;r.italic=italic
    if color:r.font.color.rgb=RGBColor.from_string(color)
def read_rows():return {r['image']:r for r in csv.DictReader((BASE/'metrics.csv').open(encoding='utf-8'))}
def authenticated(name,rows):
    row=rows[name];img=np.asarray(Image.open(BASE/row['authenticated_png']).convert('RGB'),dtype=np.uint8)
    return row,img

def prepare_pixel(rows):
    row,img=authenticated('couple',rows);iid=bytes.fromhex(row['image_id_hex']);blob=(BASE/row['layout_sidecar']).read_bytes()
    clean=exact.verify_tpe_ciphertext(img,blob,KEY,iid)
    if not clean.accepted:raise AssertionError('clean Couple control failed')
    channel=1;nblocks=(img.shape[0]//BLOCK)*(img.shape[1]//BLOCK);ids=common._group_order(KEY,channel,nblocks)[0]
    bid=int(ids[0]);by,bx=divmod(bid,img.shape[1]//BLOCK)
    tampered=img.copy();py,px=by*BLOCK,bx*BLOCK;tampered[py,px,channel]^=1
    check=exact.verify_tpe_ciphertext(tampered,blob,KEY,iid)
    if check.accepted or check.failed_groups!=((channel,bid),):raise AssertionError(f'one-pixel mismatch: {check.failed_groups}')
    clean_path=DEMO/'couple_clean_for_pixel_check.png';tamper_path=DEMO/'couple_one_pixel_tampered.png'
    Image.fromarray(img).save(clean_path);Image.fromarray(tampered).save(tamper_path)
    # 16x16 context around changed pixel enlarged with nearest-neighbour only.
    y0=max(0,py-8);x0=max(0,px-8);y1=min(img.shape[0],py+8);x1=min(img.shape[1],px+8)
    for name,a in [('couple_clean_pixel_crop.png',img),('couple_tampered_pixel_crop.png',tampered)]:
        Image.fromarray(a[y0:y1,x0:x1]).resize((256,256),Image.Resampling.NEAREST).save(DEMO/name)
    return {'image':img,'tampered':tampered,'block':bid,'channel':channel,'pixel_yx':(py,px),'failed':check.failed_groups}

def prepare_splices(rows):
    airrow,air=authenticated('airplane',rows);lenarow,lena=authenticated('lena',rows)
    for name,row,img in [('airplane',airrow,air),('lena',lenarow,lena)]:
        v=exact.verify_tpe_ciphertext(img,(BASE/row['layout_sidecar']).read_bytes(),KEY,bytes.fromhex(row['image_id_hex']))
        if not v.accepted:raise AssertionError(f'clean {name} control failed')
    nblocks=(air.shape[0]//BLOCK)*(air.shape[1]//BLOCK)
    cases=[('Lena block → Airplane','airplane','lena',air,lena,airrow,1,5,'airplane_with_lena_block_y1_x5.png'),
           ('Airplane block → Lena','lena','airplane',lena,air,lenarow,0,0,'lena_with_airplane_block_y0_x0.png')]
    results=[]
    for label,recipient_name,donor_name,recipient,donor,row,by,bx,filename in cases:
        out=recipient.copy();y0,x0=by*BLOCK,bx*BLOCK
        out[y0:y0+BLOCK,x0:x0+BLOCK]=donor[y0:y0+BLOCK,x0:x0+BLOCK]
        bid=by*(recipient.shape[1]//BLOCK)+bx;expected=set()
        for c in range(3):
            for group in common._group_order(KEY,c,nblocks):
                if bid in group:expected.add((c,int(group[0])))
        v=exact.verify_tpe_ciphertext(out,(BASE/row['layout_sidecar']).read_bytes(),KEY,bytes.fromhex(row['image_id_hex']))
        if v.accepted or set(v.failed_groups)!=expected:raise AssertionError(f'{label}: mismatch, expected={expected}; got={v.failed_groups}')
        Image.fromarray(out).save(DEMO/filename)
        results.append({'label':label,'recipient':recipient_name,'donor':donor_name,'recipient_image':recipient,
                        'donor_image':donor,'tampered_image':out,'block_yx':(by,bx),'block_id':bid,
                        'groups':tuple(sorted(expected)),'filename':filename,'accepted':v.accepted})
    with (DEMO/'splice_results_bidirectional.csv').open(newline='',encoding='utf-8') as f:csvrows=list(csv.DictReader(f))
    if len(csvrows)!=2 or any(r['accepted']!='False' or r['localization_exact']!='True' for r in csvrows):
        raise AssertionError('saved bidirectional CSV does not confirm both rejections and exact localizations')
    # Full-image originals/ciphertexts and zoom crops for both directions.
    for case in results:
        by,bx=case['block_yx'];sl=(slice(by*BLOCK,(by+1)*BLOCK),slice(bx*BLOCK,(bx+1)*BLOCK))
        for suffix,key in [('recipient_zoom','recipient_image'),('donor_zoom','donor_image'),('tampered_zoom','tampered_image')]:
            crop=case[key][sl[0],sl[1]].copy()
            Image.fromarray(crop).resize((256,256),Image.Resampling.NEAREST).save(DEMO/f"{case['recipient']}_{suffix}.png")
    return results

def add_full_image_row(doc,labels,images,width=2.05):
    table=doc.add_table(rows=2,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for i,(label,img) in enumerate(zip(labels,images)):
        c=table.cell(0,i);c.width=Inches(width);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(1)
        p.add_run().add_picture(str(img),width=Inches(width-0.08))
        c=table.cell(1,i);c.width=Inches(width);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(1)
        r=p.add_run(label);style_run(r,8,True)
        for cc in [table.cell(0,i),table.cell(1,i)]:margins(cc)
    return table

def add_crop_row(doc,files,labels,width=1.35):
    table=doc.add_table(rows=2,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for i,(file,label) in enumerate(zip(files,labels)):
        c=table.cell(0,i);c.width=Inches(width);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
        p.add_run().add_picture(str(file),width=Inches(width-0.06))
        c=table.cell(1,i);c.width=Inches(width);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
        r=p.add_run(label);style_run(r,7,False)
        for cc in [table.cell(0,i),table.cell(1,i)]:margins(cc)

def heading(doc,text,size=12):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text);style_run(r,size,True,BLUE);return p

def create_doc(rows,pixel,splices):
    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.4);sec.bottom_margin=Inches(.38);sec.left_margin=Inches(.42);sec.right_margin=Inches(.42)
    normal=doc.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(8.5);normal.font.color.rgb=RGBColor(35,45,55);normal.paragraph_format.space_after=Pt(2)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(1)
    r=p.add_run('Authentication Tamper Demonstration');style_run(r,18,True,BLUE)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    r=p.add_run('Full-size images · both Airplane/Lena splice directions · one-pixel tamper');style_run(r,9,False,None,True)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    r=p.add_run('How it works. ');r.bold=True;p.add_run('For each channel, an HMAC-SHA256 tag authenticates each DRBG-selected four-block group. Verification extracts the embedded tag, recomputes it from the received group and image identifier, rejects a mismatch, and withholds recovery.')
    for case in splices:
        recipient=case['recipient'];donor=case['donor']
        heading(doc,f"{case['label']} — full 512×512 encrypted images")
        add_full_image_row(doc,[f'{recipient.title()} clean authenticated image',f'{donor.title()} donor authenticated image',f'{recipient.title()} after 32×32 block splice'],
                           [BASE/rows[recipient]['authenticated_png'],BASE/rows[donor]['authenticated_png'],DEMO/case['filename']],width=2.08)
        groups='; '.join(f'channel {c}: group {g}' for c,g in case['groups'])
        p=doc.add_paragraph();p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(1)
        p.add_run(f"Result: ").bold=True;p.add_run(f"replaced block row {case['block_yx'][0]}, column {case['block_yx'][1]}; rejected; exact affected groups: {groups}. Recovery withheld.")
    # Full two directions are shown on page one; details and pixel test on page two.
    doc.add_page_break()
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    r=p.add_run('Close-ups and one-pixel test');style_run(r,16,True,BLUE)
    for case in splices:
        heading(doc,f"{case['recipient'].title()} recipient / {case['donor'].title()} donor block detail",10)
        files=[DEMO/f"{case['recipient']}_recipient_zoom.png",DEMO/f"{case['recipient']}_donor_zoom.png",DEMO/f"{case['recipient']}_tampered_zoom.png"]
        add_crop_row(doc,files,['Recipient block before','Donor block','Recipient block after splice'],width=1.5)
    heading(doc,'Single-pixel change: authenticated Couple image',11)
    # Show entire original and altered image plus a pixel-context enlargement.
    add_full_image_row(doc,['Clean authenticated Couple','Same image after 1-pixel edit','16× enlarged pixel neighborhood'],
                       [DEMO/'couple_clean_for_pixel_check.png',DEMO/'couple_one_pixel_tampered.png',DEMO/'couple_tampered_pixel_crop.png'],width=2.08)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
    p.add_run(f"Result: ").bold=True;p.add_run(f"one pixel at (y={pixel['pixel_yx'][0]}, x={pixel['pixel_yx'][1]}), channel {pixel['channel']}, was flipped. Verification rejected the image and localized the failure to group {pixel['block']}; recovery withheld.")
    t=doc.add_table(rows=1,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Light Shading Accent 1'
    for i,x in enumerate(['Test','Verification','Localization']):
        c=t.cell(0,i);c.text=x;shade(c,BLUE)
        for r in c.paragraphs[0].runs:style_run(r,8,True,'FFFFFF')
    rows_table=[('Lena block → Airplane','Rejected; recovery withheld','3 exact channel groups'),('Airplane block → Lena','Rejected; recovery withheld','3 exact channel groups'),('One-pixel edit','Rejected; recovery withheld','1 exact four-block group')]
    for vals in rows_table:
        cells=t.add_row().cells
        for i,text in enumerate(vals):cells[i].text=text;margins(cells[i]);
        for c in cells:
            for p in c.paragraphs:
                for r in p.runs:style_run(r,8)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(0)
    r=p.add_run('Controls: clean Airplane, clean Lena, and clean Couple outputs verified before modifications. Tests use the saved research outputs and reproducibility fixture key; this is a demonstration, not a security certification.');style_run(r,8,False,'5F6973',True)
    out=DEMO/'authentication_tamper_demo.docx';doc.save(out);print('saved',out)

def main():
    DEMO.mkdir(parents=True,exist_ok=True);rows=read_rows();pixel=prepare_pixel(rows);splices=prepare_splices(rows);create_doc(rows,pixel,splices)
if __name__=='__main__':main()
