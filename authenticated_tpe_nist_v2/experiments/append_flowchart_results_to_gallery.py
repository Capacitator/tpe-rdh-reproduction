#!/usr/bin/env python3
"""Append measured flowchart-hybrid results to the current Word gallery."""
from __future__ import annotations
import csv,shutil
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output'/'metrics'/'flowchart_hybrid'
IMAGES=OUT
CSV=OUT/'metrics.csv'
CAPACITY=OUT/'capacity_audit.csv'
DOC=ROOT/'output'/'metrics'/'block32_corrected_images'/'professor_32x32_block_effect_gallery.docx'
BACKUP=DOC.with_name('professor_32x32_block_effect_gallery_before_flowchart_update.docx')
MARKER='Flowchart-combined authenticated TPE results'

def add_cell(cell,text,bold=False):
 cell.text='';r=cell.paragraphs[0].add_run(str(text));r.bold=bold;r.font.name='Arial';r.font.size=Pt(7.5)

def main():
 if not DOC.is_file():raise FileNotFoundError(DOC)
 rows=list(csv.DictReader(CSV.open(newline='',encoding='utf-8')))
 caprows=list(csv.DictReader(CAPACITY.open(newline='',encoding='utf-8')))
 if len(rows)!=6:raise RuntimeError(f'Expected six rows, found {len(rows)}')
 doc=Document(str(DOC))
 if any(MARKER in p.text for p in doc.paragraphs):
  # Re-running this generator replaces only its own appendix, leaving the
  # earlier gallery pages untouched.
  body=doc._body._element
  found=False
  removed_pagebreak=None
  for child in list(body):
   if child.tag==qn('w:p') and MARKER in ''.join(child.itertext()):
    found=True
    # The old appendix inserted a page-break paragraph immediately before
    # its heading; remove it as well so only one break starts the replacement.
    prev=child.getprevious()
    if prev is not None and prev.tag==qn('w:p') and any(
     br.get(qn('w:type'))=='page' for br in prev.iter(qn('w:br'))
    ): removed_pagebreak=prev
   if found and child.tag!=qn('w:sectPr'): body.remove(child)
  if removed_pagebreak is not None: body.remove(removed_pagebreak)
 else:
  shutil.copy2(DOC,BACKUP)
 doc.add_page_break()
 title=doc.add_heading(MARKER,0);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
 intro=doc.add_paragraph(
  'Following the professor’s flowchart, this supplement preserves the original 32×32 block-TPE stage and adds HMAC-SHA256 authentication with HMAC_DRBG grouping and reversible RCM marking. The final image remains an encrypted block-TPE output. The legacy pair-sum-preserving substitution is applied after RCM as an outer reversible step; verification first reverses that step, then authenticates, before plaintext recovery. This ordering adjustment is disclosed below and in the implementation notes.'
 );intro.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.add_heading('PSNR and thumbnail-fidelity results',level=1)
 doc.add_paragraph(
  'Full RGB PSNR compares original pixels with the encrypted image and is not the thumbnail-preservation score. “Block-thumbnail PSNR” compares rounded RGB means for corresponding 32×32 blocks in the original and encrypted images. That is the relevant block-thumbnail fidelity measure. Authentication marking changes some block means, so both pre-authentication and final authenticated values are shown.'
 )
 table=doc.add_table(rows=1,cols=5);table.style='Table Grid';table.autofit=False
 widths=[Inches(.85),Inches(1.35),Inches(1.8),Inches(1.65),Inches(1.2)]
 headers=['Image','Full RGB PSNR dB\nTPE → Auth','32×32 thumbnail PSNR dB\nTPE → Auth','32×32 thumbnail SSIM\nTPE → Auth','Auth mode']
 for i,x in enumerate(headers):table.columns[i].width=widths[i];add_cell(table.rows[0].cells[i],x,True)
 for r in rows:
  vals=[r['image'].capitalize(),f"{float(r['base_ciphertext_full_rgb_psnr_db']):.4f} → {float(r['authenticated_ciphertext_full_rgb_psnr_db']):.4f}",f"{float(r['base_block_thumbnail_psnr_db']):.4f} → {float(r['authenticated_block_thumbnail_psnr_db']):.4f}",f"{float(r['base_block_thumbnail_ssim']):.6f} → {float(r['authenticated_block_thumbnail_ssim']):.6f}",r['auth_mode']]
  cells=table.add_row().cells
  for i,x in enumerate(vals):cells[i].width=widths[i];add_cell(cells[i],x)
 base_thumb=sum(float(r['base_block_thumbnail_psnr_db']) for r in rows)/len(rows)
 auth_thumb=sum(float(r['authenticated_block_thumbnail_psnr_db']) for r in rows)/len(rows)
 base_full=sum(float(r['base_ciphertext_full_rgb_psnr_db']) for r in rows)/len(rows)
 auth_full=sum(float(r['authenticated_ciphertext_full_rgb_psnr_db']) for r in rows)/len(rows)
 doc.add_paragraph(f"Six-image mean: full RGB PSNR {base_full:.4f} → {auth_full:.4f} dB; block-thumbnail PSNR {base_thumb:.4f} → {auth_thumb:.4f} dB.")
 doc.add_heading('Measured RCM capacity and fallback',level=2)
 doc.add_paragraph('A 256-bit group tag needs at least 256 net eligible pair slots (T + O − N) within that group. HMAC_DRBG still defines the deterministic group ordering; the source method falls back atomically to one whole-image tag if any group cannot carry its tag. The table reports every group’s capacity outcome count, not a selected subset.')
 ctable=doc.add_table(rows=1,cols=4);ctable.style='Table Grid';ctable.autofit=False
 cwidth=[Inches(.9),Inches(1.55),Inches(1.55),Inches(2.0)]
 for i,x in enumerate(['Image','Post-auth-Step-2 whole net bits','4-block groups with ≥256 bits','Legacy post-substitution net bits']):
  ctable.columns[i].width=cwidth[i];add_cell(ctable.rows[0].cells[i],x,True)
 for r in caprows:
  vals=[r['image'].capitalize(),r['post_step2_whole_net_capacity_bits'],f"{r['groups_with_capacity_ge_256']}/{r['groups_total']}",r['legacy_after_substitution_whole_net_capacity_bits']]
  cells=ctable.add_row().cells
  for i,x in enumerate(vals):cells[i].width=cwidth[i];add_cell(cells[i],x)
 doc.add_paragraph(
  'All six outputs authenticated and recovered the original image and payload exactly in the test. All six selected whole-image authentication fallback because at least one four-block group was short of 256 net RCM slots (e.g., Couple: only 22 of 48 groups fit; remaining 26 do not). After the legacy substitution, RCM net capacity is negative for every image (for example, Baboon: −133,026 bits), so embedding RCM after that legacy stage is impossible under this code. The hybrid therefore places authentication before that outer reversible substitution and inverts it before verifying. Authentication works, but group-level tamper localization is not achieved for these six images. No tag or failure was hidden.'
 )
 doc.add_page_break()
 doc.add_heading('Authenticated output examples',level=1)
 for idx,r in enumerate(rows):
  if idx in (0,2,4):
   if idx>0: doc.add_page_break()
  heading=doc.add_heading(f"{r['image'].capitalize()} — original, block TPE, and authenticated output",level=2)
  heading.paragraph_format.keep_with_next=True
  heading.paragraph_format.space_after=Pt(1)
  m=doc.add_paragraph(
   f"Full RGB PSNR: {float(r['authenticated_ciphertext_full_rgb_psnr_db']):.4f} dB | "
   f"32×32 thumbnail PSNR: {float(r['authenticated_block_thumbnail_psnr_db']):.4f} dB | "
   f"Thumbnail SSIM: {float(r['authenticated_block_thumbnail_ssim']):.6f} | "
   f"Auth: {r['auth_mode']} | Exact recovery and payload: verified"
  );m.paragraph_format.keep_with_next=True
  m.paragraph_format.space_after=Pt(1)
  t=doc.add_table(rows=2,cols=3);t.autofit=False
  for j,label in enumerate(['Original','Block TPE (before authentication)','Final authenticated TPE']):
   t.columns[j].width=Inches(2.35)
   cell=t.cell(0,j);cell.text=label;cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
   cell.paragraphs[0].paragraph_format.space_after=Pt(0)
   pic=t.cell(1,j).paragraphs[0];pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
   pic.paragraph_format.space_after=Pt(0)
   filename=[r['original_png'],r['base_ciphertext_png'],r['authenticated_ciphertext_png']][j]
   pic.add_run().add_picture(str(IMAGES/filename),width=Inches(2.05))
 note=doc.add_paragraph(
  'Limitation and next step: This is a reproducible research prototype following the flowchart’s component order at a reversible boundary. It demonstrates authenticated recovery while retaining visible block effects, but does not obtain four-block localization on this image set because RCM capacity is inadequate. To require per-group localization, the RCM carrier/capacity scheme must be redesigned and then revalidated; changing the key, block membership, or reporting convention to make this run look successful would not be scientifically appropriate.'
 )
 note.paragraph_format.space_before=Pt(2)
 doc.save(DOC)
 print(f'Updated current gallery in place: {DOC}')
 print(f'Previous version preserved: {BACKUP}')
 print(f'Appended all {len(rows)} PSNR rows and six authenticated image triplets.')
if __name__=='__main__':main()
