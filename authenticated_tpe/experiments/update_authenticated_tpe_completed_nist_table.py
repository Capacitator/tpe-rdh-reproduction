#!/usr/bin/env python3
"""Replace Table 7 in the professor's existing DOCX with actual STS rows."""
from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

TEST_LABELS = {
    "Frequency": "Frequency",
    "BlockFrequency": "Block Frequency",
    "CumulativeSums": "Cumulative Sums",
    "Runs": "Runs",
    "LongestRun": "Longest Run of Ones",
    "Rank": "Binary Matrix Rank",
    "FFT": "Discrete Fourier Transform",
    "NonOverlappingTemplate": "Non-overlapping Template Matching",
    "OverlappingTemplate": "Overlapping Template Matching",
    "Universal": "Maurer's Universal Statistical",
    "ApproximateEntropy": "Approximate Entropy",
    "RandomExcursions": "Random Excursions",
    "RandomExcursionsVariant": "Random Excursions Variant",
    "Serial": "Serial",
    "LinearComplexity": "Linear Complexity",
}
IMAGE_CODES = {
    "airplane": "air", "baboon": "bab", "couple": "cou", "girl": "girl",
    "lena": "lena", "peppers": "pep",
}
CATEGORY_ORDER = ("encrypted", "decrypted")
HEADERS = (
    "NIST Test",
    "Encrypted image p-value",
    "Encrypted result",
    "Decrypted image p-value",
    "Decrypted result",
)


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def result_cell(rows: list[dict[str, str]]) -> str:
    return "\n".join(f"{IMAGE_CODES[r['image']]} {r['first_level_result'].upper()}" for r in rows)


def pvalue_cell(rows: list[dict[str, str]]) -> str:
    return "\n".join(f"{IMAGE_CODES[r['image']]} {r['first_level_p_value']}" for r in rows)


def set_cell_text(cell, text: str, *, size: float, bold: bool = False, center: bool = False) -> None:
    tc = cell._tc
    old_p = cell.paragraphs[0]._p if cell.paragraphs else None
    p = OxmlElement("w:p")
    if old_p is not None and old_p.pPr is not None:
        p.append(deepcopy(old_p.pPr))
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(p)
    paragraph = Paragraph(p, cell)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(size + 1)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def set_row_cant_split(row) -> None:
    props = row._tr.get_or_add_trPr()
    if props.find(qn("w:cantSplit")) is None:
        props.append(OxmlElement("w:cantSplit"))


def set_table_widths(table, widths: tuple[float, ...]) -> None:
    table.autofit = False
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(round(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcw = cell._tc.tcPr.find(qn("w:tcW"))
            if tcw is not None:
                tcw.set(qn("w:w"), str(round(width * 1440)))
                tcw.set(qn("w:type"), "dxa")


def add_top_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    top = borders.find(qn("w:top"))
    if top is None:
        top = OxmlElement("w:top")
        borders.append(top)
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "8")
    top.set(qn("w:space"), "0")
    top.set(qn("w:color"), "000000")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--results", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--uniformity", type=Path, required=True)
    args = parser.parse_args()

    rows = read_rows(args.results)
    summary = __import__("json").loads(args.summary.read_text(encoding="utf-8"))
    uniformity = read_rows(args.uniformity)
    groups: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        groups[(row["test"], row["component"], row["category"])].append(row)
    ordered = list(dict.fromkeys((r["test"], r["component"]) for r in rows))
    if len(ordered) != 188:
        raise ValueError(f"Expected all 188 STS components; found {len(ordered)}")
    if len(rows) != 2 * len(ordered) * 6:
        raise ValueError(f"Expected 6 stream rows per component/category; found {len(rows)}")

    doc = Document(args.docx)
    if len(doc.tables) != 8:
        raise ValueError(f"Expected the existing eight tables; found {len(doc.tables)}")
    table = doc.tables[7]
    if len(table.columns) != 5:
        raise ValueError(f"Expected a five-column NIST table; found {len(table.columns)}")

    old_rows = table.rows
    if len(old_rows) < 3:
        raise ValueError("Existing NIST table does not have its expected header and sample rows")
    body_template = deepcopy(old_rows[2]._tr)
    table_xml = table._tbl
    note_prefixes = (
        "The specification defines encrypted image E as Step 1 followed by Step 2;",
        "NIST SP 800-22 was applied separately to the actual encrypted image output",
    )
    for paragraph in list(doc.paragraphs):
        if paragraph.text.startswith(note_prefixes):
            paragraph._element.getparent().remove(paragraph._element)
    for tr in list(table_xml.tr_lst):
        table_xml.remove(tr)
    table_xml.append(deepcopy(body_template))
    header = table.rows[0]
    for cell, label in zip(header.cells, HEADERS):
        set_cell_text(cell, "Enc. result" if label == "Encrypted result" else ("Dec. result" if label == "Decrypted result" else label), size=8, bold=True, center=True)
        add_top_border(cell)
    repeat = OxmlElement("w:tblHeader")
    header._tr.get_or_add_trPr().append(repeat)

    for test, component in ordered:
        body_xml = deepcopy(body_template)
        table_xml.append(body_xml)
        row = table.rows[-1]
        label = TEST_LABELS[test]
        if component == "single":
            visible_label = label
        elif component.startswith("m=9 template="):
            template_index = next(i for i, item in enumerate(
                [r["component"] for r in rows if r["test"] == test and r["category"] == "encrypted"]
            ) if item == component) + 1
            visible_label = f"Non-overlap template {template_index:03d}"
        else:
            visible_label = f"{label} ({component})"
        encrypted = sorted(groups[(test, component, "encrypted")], key=lambda r: int(r["stream_index"]))
        decrypted = sorted(groups[(test, component, "decrypted")], key=lambda r: int(r["stream_index"]))
        contents = (visible_label, pvalue_cell(encrypted), result_cell(encrypted), pvalue_cell(decrypted), result_cell(decrypted))
        for idx, (cell, text) in enumerate(zip(row.cells, contents)):
            set_cell_text(cell, text, size=6.2 if idx in (1, 3) else (6.5 if idx in (2, 4) else 7.0), center=idx in (1, 2, 3, 4))
        set_row_cant_split(row)

    set_table_widths(table, (1.6, 1.35, 0.9, 1.35, 0.9))

    # Retain the existing Table 7 caption position and paragraph style.
    caption = next((p for p in doc.paragraphs if p.text.startswith("Table 7.")), None)
    if caption is None:
        raise ValueError("Could not find the existing Table 7 caption")
    caption.clear()
    caption_run = caption.add_run(
        "Table 7. NIST SP 800-22 first-level p-values from STS 2.1.2 (α = 0.01). "
        "Rows show every STS component and all six image outputs; image codes are air=airplane, "
        "bab=baboon, cou=couple, girl, lena, and pep=peppers. P/F/N/A are per-stream results. "
        "Values below 1e-300 are marked <1e-300; Runs estimator-criterion zeros are identified separately."
    )
    caption_run.font.name = "Times New Roman"
    caption_run.font.size = Pt(10)

    clarification = (
        "NIST SP 800-22 was applied separately to the actual encrypted image output and the recovered decrypted image output. Authentication tags and other authentication information were excluded from this comparison. Exact recovery quality was evaluated separately using MSE, PSNR, and SSIM."
    )
    encrypted = summary["counts"]["encrypted"]
    decrypted = summary["counts"]["decrypted"]
    uniformity_states = {row["uniformity_status"] for row in uniformity}
    if uniformity_states != {"not_computed_by_STS (fewer than 55 streams)"}:
        raise ValueError(f"Unexpected second-level uniformity status: {uniformity_states}")
    notes = (
        clarification,
        f"The specification defines encrypted image E as Step 1 followed by Step 2; RCM marking occurs afterward. Very small p-values are consistent with these image outputs retaining structure: Step 1 preserves adjacent-pair sums, and the decrypted outputs are the original images. NIST SP 800-22 is not a cryptographic security test. Each category has one 1,000,000-bit MSB-first prefix from each of the six fixed UCT images. Each recovery was verified as bit-exact. These six image samples are not claimed to be independent random streams. At α = 0.01, the component-stream totals are encrypted: {encrypted['pass']} pass, {encrypted['fail']} fail, {encrypted['N/A']} N/A; decrypted: {decrypted['pass']} pass, {decrypted['fail']} fail, {decrypted['N/A']} N/A. STS reports second-level uniformity as ---- for every component because six streams are below its minimum of 55; those values are N/A, not p = 0. The complete raw per-image p-values, including all 148 non-overlapping-template components, are in first_level_p_values.csv in the authenticated-TPE evaluation outputs."
    )
    for text in notes:
        p_xml = OxmlElement("w:p")
        table_xml.addnext(p_xml)
        paragraph = Paragraph(p_xml, table._parent)
        paragraph.style = doc.styles["Normal"]
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)

    doc.save(args.docx)
    print(f"updated {args.docx}; tables={len(doc.tables)}; NIST component rows={len(ordered)}")


if __name__ == "__main__":
    main()
